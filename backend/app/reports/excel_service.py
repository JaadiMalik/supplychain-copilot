import json

from openpyxl import Workbook
from openpyxl.styles import Font

from app.reports.common import answer_text, report_path, trace_rows


def build_excel(analysis: dict):
    path = report_path(analysis["analysis_id"], "xlsx")
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Field", "Value"])
    summary["A1"].font = Font(bold=True)
    summary["B1"].font = Font(bold=True)
    summary.append(["Analysis ID", analysis["analysis_id"]])
    summary.append(["Created", analysis.get("created_at", "")])
    summary.append(["Question", analysis.get("question", "")])
    summary.append(["Answer", answer_text(analysis)])

    steps = workbook.create_sheet("Agent Steps")
    steps.append(["Round", "Tool", "Status", "Arguments", "Planner Reason"])
    for cell in steps[1]:
        cell.font = Font(bold=True)
    for item in trace_rows(analysis):
        steps.append([
            item["round"],
            item["tool"],
            item["status"],
            item["arguments"],
            item["planner_reason"],
        ])

    context_sheet = workbook.create_sheet("Context")
    context_sheet.append(["Context JSON"])
    context_sheet["A1"].font = Font(bold=True)
    context_sheet.append([json.dumps(analysis.get("context", {}), indent=2, default=str)])

    raw = workbook.create_sheet("Raw JSON")
    raw.append(["Analysis JSON"])
    raw["A1"].font = Font(bold=True)
    raw.append([json.dumps(analysis, indent=2, default=str)])

    workbook.save(path)
    return path
