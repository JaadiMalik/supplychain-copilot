from docx import Document

from app.reports.common import answer_text, report_path, trace_rows


def build_docx(analysis: dict):
    path = report_path(analysis["analysis_id"], "docx")
    document = Document()
    document.add_heading("SupplyChain Copilot Analysis", level=0)
    document.add_paragraph(f"Analysis ID: {analysis['analysis_id']}")
    document.add_paragraph(f"Created: {analysis.get('created_at', '')}")

    document.add_heading("Question", level=1)
    document.add_paragraph(analysis.get("question", ""))

    document.add_heading("Answer", level=1)
    document.add_paragraph(answer_text(analysis) or "No answer.")

    document.add_heading("Agent Steps", level=1)
    rows = trace_rows(analysis)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text = "Round"
    headers[1].text = "Tool"
    headers[2].text = "Status"
    headers[3].text = "Reason"
    for item in rows:
        cells = table.add_row().cells
        cells[0].text = str(item["round"] or "")
        cells[1].text = str(item["tool"] or "")
        cells[2].text = str(item["status"] or "")
        cells[3].text = str(item["planner_reason"] or "")

    document.save(path)
    return path
